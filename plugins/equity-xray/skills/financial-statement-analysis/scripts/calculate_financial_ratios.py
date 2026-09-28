#!/usr/bin/env python3
"""Normalize, calculate and trace nonfinancial-company statements. No network calls.
Missing inputs remain null. Audit checks are numerical, not business risk ratings.
"""
import argparse
import csv
import json
import math
from datetime import date
from pathlib import Path
from dupont_analysis import shapley


def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def div(a, b):
    return a / b if number(a) and number(b) and b > 0 else None


def mean(a, b):
    return (a + b) / 2 if number(a) and number(b) else None


def product(values):
    return math.prod(values) if all(number(x) for x in values) else None


LABELS = {
    'gross_profit': '毛利润', 'gross_margin': '毛利率', 'net_margin': '合并净利率',
    'parent_margin': '归母净利率', 'adjusted_parent_margin': '扣非归母净利率',
    'asset_turnover': '总资产周转率（期间）', 'equity_multiplier': '权益乘数（总权益）',
    'parent_equity_multiplier': '权益乘数（归母口径）', 'roe_total': '自算ROE（总权益）',
    'roe_parent': '自算ROE（归母）', 'dupont_total': '杜邦ROE（总权益）',
    'dupont_parent': '杜邦ROE（归母）', 'dso': '应收周转天数', 'dio': '存货周转天数',
    'dpo': '应付周转天数', 'ccc': '现金转换周期', 'debt_ratio': '资产负债率',
    'net_debt': '净债务', 'net_cash': '净现金', 'roa_net_profit': 'ROA（净利润口径）', 'simple_fcf': '简化自由现金流', 'cash_conversion': 'OCF/合并净利润',
    'interest_coverage': '利息覆盖（分析EBIT）', 'effective_tax_rate': '有效税率',
    'revenue_yoy': '营业收入同比', 'ar_yoy': '应收账款同比', 'inventory_yoy': '存货同比',
}


def validate(data):
    if data.get('schema_version') != 1:
        raise ValueError('schema_version must be 1')
    sources = data.get('sources', [])
    source_ids = set()
    for source in sources:
        for key in ['id', 'title', 'locator', 'published_date', 'accessed_date']:
            if not source.get(key):
                raise ValueError(f'Source missing {key}: {source.get("id")}')
        if source['id'] in source_ids:
            raise ValueError('Duplicate source id: ' + source['id'])
        source_ids.add(source['id'])
    ids = set()
    for r in data.get('records', []):
        for key in ['id', 'company', 'ticker', 'period_start', 'period_end', 'period_type', 'currency', 'unit', 'scope', 'accounting_standard', 'values']:
            if key not in r:
                raise ValueError(f'Record missing {key}')
        if r['id'] in ids:
            raise ValueError('Duplicate record id: ' + r['id'])
        ids.add(r['id'])
        if r['period_type'] not in ['FY', 'YTD', 'Q', 'TTM']:
            raise ValueError('Invalid period_type')
        if date.fromisoformat(r['period_end']) < date.fromisoformat(r['period_start']):
            raise ValueError('Invalid period dates')
        if r['scope'] != 'consolidated':
            raise ValueError('Built-in calculator requires consolidated records; handle separate statements explicitly')
        for key, cell in r['values'].items():
            if not isinstance(cell, dict) or 'value' not in cell:
                raise ValueError(f'{r["id"]}.{key} must be a sourced cell object')
            if cell['value'] is not None and not number(cell['value']):
                raise ValueError(f'Nonfinite/non-numeric value: {r["id"]}.{key}')
            if cell['value'] is not None:
                if not cell.get('source_ids') or not cell.get('locator') or not cell.get('original_label'):
                    raise ValueError(f'Missing cell provenance: {r["id"]}.{key}')
                if not set(cell['source_ids']).issubset(source_ids):
                    raise ValueError(f'Unknown source: {r["id"]}.{key}')
            if key not in ['roe_reported', 'roe_adjusted_reported'] and not key.startswith('extra_') and cell.get('unit', r['unit']) != r['unit']:
                raise ValueError('Normalize monetary cells to record unit before calculation')
            if key in ['roe_reported', 'roe_adjusted_reported'] and cell.get('unit') != 'ratio':
                raise ValueError('Reported ROE needs ratio unit, e.g. 0.15 for 15%')
        tolerance = r.get('rounding_tolerance', 0.01)
        if not number(tolerance) or tolerance < 0:
            raise ValueError('rounding_tolerance must be nonnegative')
    if not ids:
        raise ValueError('At least one record is required')
    return source_ids


def analyze(data):
    validate(data)
    raw_rows, metric_rows, checks = [], [], []
    by_record = {r['id']: r for r in data['records']}
    nodes = {}
    by_metrics = {}
    for r in data['records']:
        rid = r['id']
        for key, cell in r['values'].items():
            identifier = f'{rid}:raw:{key}'
            row = dict(id=identifier, record_id=rid, company=r['company'], ticker=r['ticker'],
                       period_start=r['period_start'], period_end=r['period_end'], period_type=r['period_type'],
                       currency=r['currency'], unit=cell.get('unit', r['unit']), field=key, value=cell['value'],
                       source_ids=cell.get('source_ids', []), locator=cell.get('locator', ''),
                       original_label=cell.get('original_label', ''), original_value=cell.get('original_value'),
                       original_unit=cell.get('original_unit', ''), note=cell.get('note', ''))
            raw_rows.append(row)
            nodes[identifier] = row

    for r in data['records']:
        rid, cells = r['id'], r['values']
        vals = {k: c['value'] for k, c in cells.items()}
        m = {}
        by_metrics[rid] = m
        days = (date.fromisoformat(r['period_end']) - date.fromisoformat(r['period_start'])).days + 1
        days = r.get('day_basis', days)
        if not number(days) or days <= 0:
            raise ValueError('day_basis must be positive')

        def raw(key):
            return vals.get(key)

        def add(key, value, dependencies, formula, unit='ratio', note=''):
            dependencies = [d if ':raw:' in d or ':metric:' in d else f'{rid}:raw:{d}' for d in dependencies]
            known = [nodes[d] for d in dependencies if d in nodes]
            source_ids = sorted({s for n in known for s in n.get('source_ids', [])})
            row = dict(id=f'{rid}:metric:{key}', record_id=rid, company=r['company'], ticker=r['ticker'],
                       period_start=r['period_start'], period_end=r['period_end'], period_type=r['period_type'],
                       metric=key, label=LABELS.get(key, key), value=value, unit=unit, formula=formula,
                       dependencies=dependencies, source_ids=source_ids,
                       status='computed' if value is not None else 'unavailable',
                       note=note if value is not None else (note + '；缺少输入或分母/口径不适用').strip('；'))
            m[key] = value
            nodes[row['id']] = row
            metric_rows.append(row)
            return value

        def mid(key):
            return f'{rid}:metric:{key}'

        def check(name, lhs, rhs, note=''):
            delta = lhs - rhs if number(lhs) and number(rhs) else None
            tol = max(r.get('rounding_tolerance', 0.01), 1e-9 * max(abs(lhs or 0), abs(rhs or 0)))
            checks.append(dict(record_id=rid, check=name, lhs=lhs, rhs=rhs, difference=delta,
                               tolerance=tol, status='NOT_TESTED' if delta is None else ('PASS' if abs(delta) <= tol else 'FAIL'), note=note))

        def arithmetic(keys, signs):
            return sum(raw(k) * sign for k, sign in zip(keys, signs)) if all(number(raw(k)) for k in keys) else None

        gp = add('gross_profit', arithmetic(['revenue', 'cogs'], [1, -1]), ['revenue', 'cogs'], 'revenue - cogs', r['unit'])
        for key, num in [('gross_margin', 'gross_profit'), ('net_margin', 'net_profit'), ('parent_margin', 'parent_profit'), ('adjusted_parent_margin', 'adjusted_parent_profit')]:
            add(key, div(gp if num == 'gross_profit' else raw(num), raw('revenue')), [mid('gross_profit') if num == 'gross_profit' else num, 'revenue'], f'{num} / revenue')
        for key in ['selling_expense', 'admin_expense', 'rd_expense', 'financial_expense', 'capex_cash']:
            add(key + '_ratio', div(raw(key), raw('revenue')), [key, 'revenue'], f'{key} / revenue')
        for key in ['total_assets', 'total_equity', 'parent_equity', 'ar', 'inventory', 'trade_payables', 'fixed_assets']:
            add('avg_' + key, mean(raw(key + '_open'), raw(key + '_close')), [key + '_open', key + '_close'], f'({key}_open + {key}_close) / 2', r['unit'])
        add('roa_net_profit', div(raw('net_profit'), m['avg_total_assets']), ['net_profit', mid('avg_total_assets')], 'consolidated_net_profit / avg_total_assets; period basis, not annualized')
        for base in ['total_equity', 'parent_equity']:
            a, b = raw(base + '_open'), raw(base + '_close')
            valid = number(a) and number(b) and min(a, b) > 0
            avg = m['avg_' + base] if valid else None
            suffix = 'total' if base == 'total_equity' else 'parent'
            profit = 'net_profit' if suffix == 'total' else 'parent_profit'
            add('roe_' + suffix, div(raw(profit), avg), [profit, mid('avg_' + base)], f'{profit} / avg_{base}; require both equity endpoints > 0')
            ek = 'equity_multiplier' if suffix == 'total' else 'parent_equity_multiplier'
            add(ek, div(m['avg_total_assets'], avg), [mid('avg_total_assets'), mid('avg_' + base)], f'avg_total_assets / avg_{base}; require both equity endpoints > 0', 'times')
        add('asset_turnover', div(raw('revenue'), m['avg_total_assets']), ['revenue', mid('avg_total_assets')], 'revenue / avg_total_assets', 'times_in_period')
        for suffix, margin, em in [('total', 'net_margin', 'equity_multiplier'), ('parent', 'parent_margin', 'parent_equity_multiplier')]:
            d = add('dupont_' + suffix, product([m[margin], m['asset_turnover'], m[em]]), [mid(margin), mid('asset_turnover'), mid(em)], f'{margin} * asset_turnover * {em}')
            check('dupont_' + suffix, d, m['roe_' + suffix], 'Analytical average-balance ROE, not reported weighted ROE')
        for key, numerator, average in [('ar_turnover', 'revenue', 'avg_ar'), ('inventory_turnover', 'cogs', 'avg_inventory'), ('fixed_asset_turnover', 'revenue', 'avg_fixed_assets')]:
            add(key, div(raw(numerator), m[average]), [numerator, mid(average)], f'{numerator} / {average}', 'times_in_period')
        for key, average, denom in [('dso', 'avg_ar', 'revenue'), ('dio', 'avg_inventory', 'cogs')]:
            ratio = div(m[average], raw(denom))
            add(key, ratio * days if ratio is not None else None, [mid(average), denom], f'{days} * {average} / {denom}', 'days')
        purchase_key = 'purchases' if number(raw('purchases')) else 'cogs'
        ratio = div(m['avg_trade_payables'], raw(purchase_key))
        add('dpo', ratio * days if ratio is not None else None, [mid('avg_trade_payables'), purchase_key], f'{days} * avg_trade_payables / {purchase_key}', 'days', '营业成本替代采购额的近似' if purchase_key == 'cogs' else '')
        add('ccc', m['dso'] + m['dio'] - m['dpo'] if all(number(m[x]) for x in ['dso', 'dio', 'dpo']) else None, [mid(x) for x in ['dso', 'dio', 'dpo']], 'dso + dio - dpo', 'days')
        add('debt_ratio', div(raw('total_liabilities_close'), raw('total_assets_close')), ['total_liabilities_close', 'total_assets_close'], 'total_liabilities_close / total_assets_close')
        add('net_debt', arithmetic(['interest_debt_close', 'available_cash_close'], [1, -1]), ['interest_debt_close', 'available_cash_close'], 'interest_debt_close - available_cash_close', r['unit'])
        add('net_cash', -m['net_debt'] if m['net_debt'] is not None else None, [mid('net_debt')], '-net_debt', r['unit'])
        add('debt_to_available_cash', div(raw('interest_debt_close'), raw('available_cash_close')), ['interest_debt_close', 'available_cash_close'], 'interest_debt_close / available_cash_close', 'times')
        ebit = add('ebit_analytical', arithmetic(['profit_before_tax', 'interest_expense'], [1, 1]), ['profit_before_tax', 'interest_expense'], 'profit_before_tax + expensed_interest; includes nonoperating items', r['unit'])
        add('interest_coverage', div(ebit, raw('interest_expense')), [mid('ebit_analytical'), 'interest_expense'], 'ebit_analytical / interest_expense', 'times')
        add('effective_tax_rate', div(raw('income_tax'), raw('profit_before_tax')), ['income_tax', 'profit_before_tax'], 'income_tax / profit_before_tax')
        add('cash_conversion', div(raw('cfo'), raw('net_profit')), ['cfo', 'net_profit'], 'cfo / consolidated_net_profit; positive denominator only', 'times', '微利分母仍需人工判断，不能据比率机械评级')
        add('simple_fcf', arithmetic(['cfo', 'capex_cash'], [1, -1]), ['cfo', 'capex_cash'], 'cfo - cash_capex; not automatically FCFF', r['unit'])
        add('ocf_profit_gap', arithmetic(['cfo', 'net_profit'], [1, -1]), ['cfo', 'net_profit'], 'cfo - consolidated_net_profit', r['unit'])
        add('parent_adjustment_gap', arithmetic(['parent_profit', 'adjusted_parent_profit'], [1, -1]), ['parent_profit', 'adjusted_parent_profit'], 'parent_profit - adjusted_parent_profit', r['unit'])
        add('goodwill_equity', div(raw('goodwill_close'), raw('total_equity_close')), ['goodwill_close', 'total_equity_close'], 'goodwill_close / total_equity_close')
        check('balance_sheet', raw('total_assets_close'), arithmetic(['total_liabilities_close', 'total_equity_close'], [1, 1]))
        check('profit_attribution', raw('net_profit'), arithmetic(['parent_profit', 'minority_profit'], [1, 1]))
        check('pretax_to_net', raw('net_profit'), arithmetic(['profit_before_tax', 'income_tax'], [1, -1]))
        check('gross_profit_disclosed', gp, raw('gross_profit_disclosed'))
        check('cash_rollforward', arithmetic(['cash_equiv_close', 'cash_equiv_open'], [1, -1]), arithmetic(['cfo', 'cfi', 'cff', 'fx_cash_change'], [1, 1, 1, 1]))
        for bridge in r.get('bridges', []):
            terms = bridge['terms']
            if not terms or any(t['sign'] not in [-1, 1] for t in terms):
                raise ValueError('Bridge terms need signs +1 or -1')
            check(bridge['name'], arithmetic([t['field'] for t in terms], [t['sign'] for t in terms]), raw(bridge['target']), bridge.get('note', ''))

    # Explicit prior references prevent accidental FY/TTM or nonadjacent comparisons.
    for r in data['records']:
        prior_id = r.get('prior_record_id')
        if not prior_id:
            continue
        if prior_id not in by_record:
            raise ValueError('Unknown prior_record_id: ' + prior_id)
        p = by_record[prior_id]
        for key in ['ticker', 'period_type', 'currency', 'unit', 'scope', 'accounting_standard']:
            if r[key] != p[key]:
                raise ValueError(f'Incomparable prior record: {key}')
        for key in ['period_start', 'period_end']:
            old, new = date.fromisoformat(p[key]), date.fromisoformat(r[key])
            if new.year - old.year != 1 or new.month != old.month or abs(new.day - old.day) > 1:
                raise ValueError('YoY prior must match prior-year fiscal window; custom calendars require explicit custom calculations')
        for balance in ['total_assets','total_equity','parent_equity','ar','inventory','trade_payables','fixed_assets','cash_equiv']:
            opening = r['values'].get(balance+'_open',{}).get('value')
            previous_close = p['values'].get(balance+'_close',{}).get('value')
            adjustment_cell = r['values'].get('opening_adjustment_'+balance)
            adjustment = adjustment_cell['value'] if adjustment_cell else 0
            rhs = previous_close + adjustment if number(previous_close) and number(adjustment) else None
            difference = opening - rhs if number(opening) and number(rhs) else None
            tolerance = max(r.get('rounding_tolerance',.01),1e-9*max(abs(opening or 0),abs(rhs or 0)))
            checks.append(dict(record_id=r['id'],check='opening_continuity_'+balance,lhs=opening,rhs=rhs,
                               difference=difference,tolerance=tolerance,
                               status='NOT_TESTED' if difference is None else ('PASS' if abs(difference)<=tolerance else 'FAIL'),
                               note='Opening versus prior close; any restatement adjustment must be a sourced opening_adjustment_* cell'))
        for key, field in [('revenue_yoy', 'revenue'), ('ar_yoy', 'ar_close'), ('inventory_yoy', 'inventory_close')]:
            a, b = p['values'].get(field, {}).get('value'), r['values'].get(field, {}).get('value')
            ratio = div(b, a) if number(b) and b >= 0 else None
            deps = [f'{p["id"]}:raw:{field}', f'{r["id"]}:raw:{field}']
            row = dict(id=f'{r["id"]}:metric:{key}', record_id=r['id'], company=r['company'], ticker=r['ticker'],
                       period_start=r['period_start'], period_end=r['period_end'], period_type=r['period_type'],
                       metric=key, label=LABELS[key], value=ratio-1 if ratio is not None else None, unit='ratio',
                       formula=f'{deps[1]} / {deps[0]} - 1', dependencies=deps,
                       source_ids=sorted({s for d in deps for s in nodes.get(d, {}).get('source_ids', [])}),
                       status='computed' if ratio is not None else 'unavailable', note='Low base requires analyst review')
            nodes[row['id']] = row
            metric_rows.append(row)
        for basis, margin, em in [('total', 'net_margin', 'equity_multiplier'), ('parent', 'parent_margin', 'parent_equity_multiplier')]:
            keys = [margin, 'asset_turnover', em]
            before = [by_metrics[p['id']][k] for k in keys]
            after = [by_metrics[r['id']][k] for k in keys]
            contributions = shapley(before, after)
            for i, k in enumerate(keys):
                deps = [f'{rid}:metric:{metric}' for rid in [p['id'], r['id']] for metric in keys]
                row = dict(id=f'{r["id"]}:metric:roe_{basis}_contribution_{k}', record_id=r['id'],
                           company=r['company'], ticker=r['ticker'], period_start=r['period_start'], period_end=r['period_end'], period_type=r['period_type'],
                           metric=f'roe_{basis}_contribution_{k}', label=f'ROE变动贡献（{basis}）: {k}',
                           value=contributions[i]*100 if contributions else None, unit='percentage_points',
                           formula='Shapley mean of marginal product changes over all 6 factor orders; multiply by 100',
                           dependencies=deps, source_ids=sorted({s for d in deps for s in nodes[d]['source_ids']}),
                           status='computed' if contributions else 'unavailable', note='Mathematical attribution, not causal proof')
                nodes[row['id']] = row
                metric_rows.append(row)
    for request in data.get('trend_requests', []):
        records = [by_record[i] for i in request['record_ids']]
        if len(records) < 2 or len({r['id'] for r in records}) != len(records):
            raise ValueError('Trend requests require at least two unique periods')
        first = records[0]
        for i, r in enumerate(records):
            if r['period_type'] != 'FY' or any(r[k] != first[k] for k in ['ticker','currency','unit','scope','accounting_standard']):
                raise ValueError('Trend requests require comparable full fiscal years')
            if i and r.get('prior_record_id') != records[i-1]['id']:
                raise ValueError('Trend requests require an explicit continuous comparable chain')
        kind = request['kind']
        if kind == 'cagr':
            field = request['field']
            deps = [f'{r["id"]}:raw:{field}' for r in records]
            values = [nodes.get(d, {}).get('value') for d in deps]
            n = len(records) - 1
            value = (values[-1]/values[0])**(1/n)-1 if all(number(x) and x>0 for x in values) else None
            formula = f'(last {field} / first {field}) ** (1/{n}) - 1'
            unit = 'ratio'
        elif kind == 'cumulative_cash_conversion':
            deps = [f'{r["id"]}:raw:{field}' for r in records for field in ['cfo','net_profit']]
            values = [nodes.get(d, {}).get('value') for d in deps]
            value = div(sum(values[::2]),sum(values[1::2])) if all(number(x) for x in values) else None
            formula = 'sum(cfo) / sum(consolidated_net_profit); positive cumulative denominator'
            unit = 'times'
        else:
            raise ValueError('Unknown trend request kind')
        identifier = 'summary:metric:' + request['id']
        if identifier in nodes:
            raise ValueError('Duplicate trend request id')
        row = dict(id=identifier, record_id='summary', company=first['company'], ticker=first['ticker'],
                   period_start=first['period_start'],period_end=records[-1]['period_end'],period_type='MULTI_FY',
                   metric=request['id'],label=request.get('label',request['id']),value=value,unit=unit,
                   formula=formula,dependencies=deps,
                   source_ids=sorted({s for d in deps for s in nodes.get(d,{}).get('source_ids',[])}),
                   status='computed' if value is not None else 'unavailable',note='Comparable continuous fiscal years; verify low-base interpretation')
        nodes[identifier]=row
        metric_rows.append(row)
    return dict(schema_version=1, metadata=data.get('metadata', {}), records=data['records'], sources=data['sources'],
                raw=raw_rows, metrics=metric_rows, checks=checks, adjustments=data.get('adjustments', []),
                peers=data.get('peers', []), claims=data.get('claims', []), accounting_notes=data.get('accounting_notes', []))


def csv_rows(path, rows, default_fields):
    fields = list(dict.fromkeys(k for row in rows for k in row)) or default_fields
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow({k: json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v for k, v in row.items()})


def export(data, result, destination):
    out = Path(destination)
    out.mkdir(parents=True, exist_ok=True)
    (out / 'input.json').write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
    (out / 'analysis.json').write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
    sets = {key: result[key] for key in ['raw', 'metrics', 'checks', 'sources', 'adjustments', 'peers', 'claims', 'accounting_notes']}
    sets['formulas'] = [{k: m[k] for k in ['id', 'label', 'value', 'unit', 'formula', 'dependencies', 'source_ids', 'note']} for m in result['metrics']]
    for key, rows in sets.items():
        if any(not isinstance(row, dict) for row in rows):
            raise ValueError(key + ' must be a list of objects')
        csv_rows(out / (key + '.csv'), rows, ['id', 'note'])
    excel = False
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        pass
    else:
        wb = Workbook()
        wb.remove(wb.active)
        for title, rows in sets.items():
            ws = wb.create_sheet(title)
            headers = list(dict.fromkeys(k for row in rows for k in row)) or ['id', 'note']
            ws.append(headers)
            for row in rows:
                ws.append([json.dumps(row.get(k), ensure_ascii=False) if isinstance(row.get(k), (dict, list)) else row.get(k) for k in headers])
            for row in ws:
                for cell in row:
                    if isinstance(cell.value, str):
                        # Source text beginning '=' must remain literal, never spreadsheet code.
                        cell.data_type = 's'
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
            for cell in ws[1]:
                cell.fill = PatternFill('solid', fgColor='183A56')
                cell.font = Font(color='FFFFFF', bold=True)
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for i, key in enumerate(headers, 1):
                ws.column_dimensions[get_column_letter(i)].width = 22 if key not in ['formula', 'dependencies', 'note', 'locator'] else 55
        wb.save(out / 'workpapers.xlsx')
        excel = True
    return excel


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input')
    parser.add_argument('--out', required=True)
    args = parser.parse_args()
    data = json.loads(Path(args.input).read_text(encoding='utf-8'))
    result = analyze(data)
    excel = export(data, result, args.out)
    counts = {s: sum(c['status'] == s for c in result['checks']) for s in ['PASS', 'FAIL', 'NOT_TESTED']}
    print(json.dumps(dict(records=len(data['records']), metrics=len(result['metrics']), checks=counts, excel=excel, out=args.out), ensure_ascii=False))
    if counts['FAIL']:
        raise SystemExit(2)


if __name__ == '__main__':
    main()
