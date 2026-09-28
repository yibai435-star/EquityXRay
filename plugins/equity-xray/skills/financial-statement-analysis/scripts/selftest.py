#!/usr/bin/env python3
"""Synthetic boundary checks. Never use these numbers as company evidence."""
import copy
import json
import math
import tempfile
from pathlib import Path
from calculate_financial_ratios import analyze, export
from dupont_analysis import shapley, analyze_dupont


def fixture():
    src = dict(id='SIM', title='明确虚构：口径验证用模拟财务表', locator='synthetic://financial-fixture', published_date='2025-03-01', accessed_date='2026-09-27')
    records=[]
    for year, revenue, assets_open, assets_close, eq_open, eq_close, net, parent, parent_open, parent_close, cfo in [
        (2023,100,90,110,40,60,10,8,30,50,15),
        (2024,120,110,210,60,100,12,12,50,100,3)]:
        raw=dict(revenue=revenue,cogs=revenue*.6,net_profit=net,parent_profit=parent,
                 adjusted_parent_profit=parent-1,minority_profit=net-parent,
                 total_assets_open=assets_open,total_assets_close=assets_close,
                 total_equity_open=eq_open,total_equity_close=eq_close,
                 parent_equity_open=parent_open,parent_equity_close=parent_close,
                 total_liabilities_close=assets_close-eq_close,
                 ar_open=10 if year==2023 else 20,ar_close=20 if year==2023 else 40,
                 inventory_open=12 if year==2023 else 18,inventory_close=18 if year==2023 else 45,
                 trade_payables_open=10 if year==2023 else 14,trade_payables_close=14,
                 profit_before_tax=net+3,income_tax=3,interest_expense=2,
                 cfo=cfo,capex_cash=7 if year==2023 else 25,cfi=-7 if year==2023 else -25,
                 cff=0 if year==2023 else 30,fx_cash_change=0,cash_equiv_open=30 if year==2023 else 38,cash_equiv_close=38 if year==2023 else 46,
                 available_cash_close=38 if year==2023 else 46,interest_debt_close=20 if year==2023 else 45,
                 financial_expense=-1,selling_expense=5,admin_expense=3,rd_expense=10,
                 gross_profit_disclosed=revenue*.4)
        record=dict(id=f'Y{year}',company='示例制造（完全虚构）',ticker='SIMULATED',period_start=f'{year}-01-01',period_end=f'{year}-12-31',period_type='FY',currency='CNY',unit='百万元',scope='consolidated',accounting_standard='模拟中国准则',values={k:dict(value=v,source_ids=['SIM'],locator='模拟表：'+k,original_label=k) for k,v in raw.items()})
        if year==2024:record['prior_record_id']='Y2023'
        records.append(record)
    return dict(schema_version=1,metadata={'synthetic':True,'purpose':'验证口径，禁止作为真实公司证据'},sources=[src],records=records,adjustments=[],peers=[],claims=[],accounting_notes=[])


def run():
    data=fixture(); result=analyze(data)
    metrics={m['id']:m['value'] for m in result['metrics']}
    get=lambda year,key:metrics[f'Y{year}:metric:{key}']
    assert math.isclose(get(2023,'roe_total'),.20)
    assert math.isclose(get(2023,'roa_net_profit'),.10)
    assert math.isclose(get(2024,'net_cash'),1.0)
    result_dupont=analyze_dupont(result)
    assert len(result_dupont['periods'])==4 and len(result_dupont['changes'])==2
    assert math.isclose(result_dupont['changes'][0]['delta_pp'],-5.0)
    assert math.isclose(get(2024,'roe_parent'),.16) # Parent denominator 75, not total equity 80.
    assert math.isclose(get(2024,'roe_total'),.15)
    assert math.isclose(get(2024,'dio'),366*31.5/72) # COGS, not revenue; leap year actual days.
    assert math.isclose(get(2024,'simple_fcf'),-22)
    assert get(2024,'financial_expense_ratio')<0
    assert math.isclose(sum(get(2024,'roe_total_contribution_'+k) for k in ['net_margin','asset_turnover','equity_multiplier']),-5)
    assert not [x for x in result['checks'] if x['status']=='FAIL']
    discontinuity=copy.deepcopy(data);discontinuity['records'][1]['values']['ar_open']['value']=10
    assert any(c['check']=='opening_continuity_ar' and c['status']=='FAIL' for c in analyze(discontinuity)['checks'])
    corrected=copy.deepcopy(discontinuity)
    corrected['records'][1]['values']['opening_adjustment_ar']=dict(value=-10,source_ids=['SIM'],locator='模拟重述说明',original_label='期初重述调整')
    assert any(c['check']=='opening_continuity_ar' and c['status']=='PASS' for c in analyze(corrected)['checks'])
    negative=copy.deepcopy(data)
    negative['records'][1]['values']['net_profit']['value']=-2
    neg=analyze(negative)
    assert next(m['value'] for m in neg['metrics'] if m['id']=='Y2024:metric:cash_conversion') is None
    assert any(c['status']=='FAIL' for c in neg['checks'])
    missing=copy.deepcopy(data);del missing['records'][1]['values']['total_assets_open']
    mm={m['id']:m['value'] for m in analyze(missing)['metrics']}
    assert mm['Y2024:metric:asset_turnover'] is None
    assert mm['Y2024:metric:dupont_total'] is None
    equity=copy.deepcopy(data);equity['records'][1]['values']['total_equity_open']['value']=-10
    em={m['id']:m['value'] for m in analyze(equity)['metrics']}
    assert em['Y2024:metric:roe_total'] is None
    mix=copy.deepcopy(data);mix['records'][1]['currency']='USD'
    try:analyze(mix)
    except ValueError:pass
    else:raise AssertionError('Mixed currencies should reject YoY')
    no_source=copy.deepcopy(data);no_source['records'][0]['values']['revenue']['source_ids']=[]
    try:analyze(no_source)
    except ValueError:pass
    else:raise AssertionError('Unsourced number accepted')
    for before,after in [([.1,.7,1.5],[.08,.9,2]),([-.1,.5,2],[.1,.5,2]),([0,1,1],[.1,2,3])]:
        assert math.isclose(sum(shapley(before,after)),math.prod(after)-math.prod(before),abs_tol=1e-12)
    with tempfile.TemporaryDirectory() as tmp:
        excel=export(data,result,tmp)
        assert json.loads((Path(tmp)/'analysis.json').read_text())['metrics']
        assert (Path(tmp)/'raw.csv').read_bytes().startswith(b'\xef\xbb\xbf')
        if excel:
            from openpyxl import load_workbook
            wb=load_workbook(Path(tmp)/'workpapers.xlsx')
            assert {'raw','metrics','formulas','sources','checks'}.issubset(wb.sheetnames)
    print('PASS: parent/consolidated ROE, Shapley, missing values, losses, negative equity, period/currency, COGS denominator, signed costs, provenance, checks, CSV/XLSX export')


if __name__=='__main__':run()
